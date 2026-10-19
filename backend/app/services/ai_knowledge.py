"""Section-level DOCX ingestion and BM25 retrieval; no embedding service."""
import re
import threading
from pathlib import Path
from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_community.retrievers import BM25Retriever

STOP_WORDS = set('a an the is are do does can could would should i you we my your me our us it this that to for of in on and or with about what how where who which please'.split())
ALIASES = {'located': 'location', 'based': 'location', 'office': 'location',
           'cost': 'pricing', 'costs': 'pricing', 'price': 'pricing', 'prices': 'pricing',
           'charge': 'pricing', 'charges': 'pricing', 'fee': 'pricing', 'fees': 'pricing',
           # Visitors describe project cost as investment, spend or budget.
           # This normalises search vocabulary, never supplies a price or quote.
           'investment': 'pricing', 'invest': 'pricing', 'budget': 'pricing',
           'budgets': 'pricing', 'spend': 'pricing', 'spending': 'pricing',
           'restaurant': 'restaurants', 'cafe': 'cafes', 'reel': 'reels',
           'instagram': 'social', 'facebook': 'meta', 'advertising': 'ads',
           'website': 'websites', 'founders': 'founder', 'duration': 'timeline',
           'company': 'agency', 'founded': 'founder', 'founding': 'founder',
           'logo': 'branding', 'logos': 'branding', 'clothing': 'fashion'}

def retrieval_query(text):
    # Questions without a named topic still have a clear company-overview intent.
    # Expand the search, not the answer: all facts still come from the document.
    normalized = ' '.join(re.findall(r'[a-z0-9]+', text.lower()))
    normalized = re.sub(r'^(hi|hello|hey)( there)?\s+', '', normalized)
    normalized = re.sub(r'\byou (guys|folks|all)\b', 'you', normalized)
    normalized = re.sub(r'\byourselves\b', 'yourself', normalized)
    normalized = re.sub(r'^what is it that (growthway|your company) does$', r'what is \1 about', normalized)
    if re.fullmatch(
        r'(what (do you do|are you|is (growthway|this company|your company|this agency)( all)?( about)?)|'
        r'who are you|tell me about (yourself|growthway|this company|your company)|'
        r'(give me )?(a |an )?(company overview|introduction))', normalized
    ):
        return 'creative growth agency identity digital presence strategy branding'
    if re.search(r'\bonline (store|stores|shop|shops)\b', normalized):
        return text + ' website ecommerce'
    return text

def tokenize(text):
    return [ALIASES.get(t, t) for t in re.findall(r'[a-z0-9]+', text.lower()) if t not in STOP_WORDS]

def read_records(path):
    word = WordDocument(path)
    paragraphs = [p.text.strip() for p in word.paragraphs if p.text.strip()]
    version = next((p for p in paragraphs if p.startswith('Retrieval reference')), '')
    records, current, body = [], None, []
    def flush():
        if current and body:
            record_id, title = current
            records.append(Document(page_content=title + '\n' + '\n'.join(body), metadata={
                'record_id': record_id, 'title': title, 'source': path.name,
                'version': version,
            }))
    for paragraph in paragraphs:
        if paragraph == 'Source register':
            break
        match = re.match(r'^((?:GW\d{3})|(?:FAQ\d{2}))\s+(.+)$', paragraph)
        if match:
            flush(); current = match.groups(); body = []
        elif current:
            if paragraph in {'Services', 'Industries and working approach', 'Portfolio',
                             'Project enquiries and contact', 'Frequently asked questions'}:
                flush(); current = None; body = []
            elif not paragraph.startswith(('Source:', 'Each answer is', 'Growthway has eight official')):
                body.append(paragraph)
    flush()
    ids = [d.metadata['record_id'] for d in records]
    if not records or len(ids) != len(set(ids)) or any(len(d.page_content) > 5000 for d in records):
        raise ValueError('Invalid company knowledge records')
    return records

class KnowledgeBase:
    def __init__(self, path, top_k=5):
        self.path = Path(path)
        self.top_k = top_k
        self._lock = threading.Lock()
        self._signature = None
        self.refresh()

    def refresh(self):
        stat = self.path.stat()
        signature = (stat.st_mtime_ns, stat.st_size)
        if signature != self._signature:
            documents = read_records(self.path)
            retriever = BM25Retriever.from_documents(documents, preprocess_func=tokenize, k=len(documents))
            self.documents, self.retriever = documents, retriever
            self._signature = signature

    def retrieve(self, query):
        with self._lock:
            self.refresh()  # Removed/invalid files fail closed, never serve stale facts.
            query = retrieval_query(query)
            tokens = tokenize(query)
            if not tokens:
                return []
            ranked = self.retriever.invoke(query)
            selected, signatures = [], []
            for document in ranked:
                words = set(tokenize(document.page_content))
                # BM25 returns k items even on zero overlap. Reject those.
                if not words.intersection(tokens):
                    continue
                if any(len(words & old) / max(1, len(words | old)) > .72 for old in signatures):
                    continue
                selected.append(document); signatures.append(words)
                if len(selected) == self.top_k:
                    break
            return selected
