"""Offline retrieval checks against the owner's actual v1.1 reference."""
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.core.config import BACKEND_DIR
from app.services.ai_knowledge import KnowledgeBase, read_records

REFERENCE = BACKEND_DIR / 'app/info/Growthway_Retrieval_Knowledge_v1.1.docx'

class RetrievalTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.kb = KnowledgeBase(REFERENCE)

    def test_record_boundaries_and_metadata(self):
        records = read_records(REFERENCE)
        self.assertEqual(len(records), 40)
        self.assertEqual(len({d.metadata['record_id'] for d in records}), 40)
        self.assertTrue(all('Version 1.1' in d.metadata['version'] for d in records))
        self.assertFalse(any('Source register' in d.page_content for d in records))
        location = next(d for d in records if d.metadata['record_id'] == 'GW004')
        self.assertIn('Lucknow', location.page_content)
        self.assertNotIn('Services\n', location.page_content)

    def test_business_questions_and_paraphrases(self):
        for query, expected in [('Where is your office located?', 'Lucknow'),
                                ('What do you charge?', 'USD'),
                                ('Can you make an Instagram reel?', 'reels'),
                                ('Who are the founders?', 'Priyanka'),
                                ('How long does a project take?', 'timeline')]:
            with self.subTest(query=query):
                text = '\n'.join(d.page_content for d in self.kb.retrieve(query))
                self.assertIn(expected.lower(), text.lower())

    def test_unknown_words_do_not_return_arbitrary_records(self):
        self.assertEqual(self.kb.retrieve('quasar thermodynamics astrophysics'), [])

    def test_project_investment_paraphrases_find_budget_evidence(self):
        for question in ['what is the investment needed?', 'what is their typical price ranges',
                         'What budget should I plan for?', 'How much should I spend?',
                         'What are your fees?', 'How much do I need to invest?']:
            with self.subTest(question=question):
                matches = self.kb.retrieve(question)
                self.assertTrue({'GW019', 'FAQ16'} & {d.metadata['record_id'] for d in matches[:2]})
                self.assertIn('fees are discussed with the team', ' '.join(d.page_content for d in matches))

    def test_company_introduction_questions_retrieve_overview(self):
        for question in ['what is this company about?', 'What is Growthway?',
                         'What do you do?', 'Who are you?', 'Tell me about your company',
                         'Give me a company overview']:
            with self.subTest(question=question):
                matches = self.kb.retrieve(question)
                self.assertTrue({'GW001', 'FAQ01'} & {d.metadata['record_id'] for d in matches[:3]})
                self.assertIn('creative growth agency', ' '.join(d.page_content.lower() for d in matches))

    def test_everyday_phrasing_finds_relevant_evidence(self):
        for question, expected in [
            ('What do you guys do?', 'GW001'),
            ('Tell me about yourselves', 'GW001'),
            ('Hi, what do you folks do?', 'GW001'),
            ('What is it that Growthway does?', 'GW001'),
            ('Who founded Growthway?', 'GW003'),
            ('Can you help with my logo?', 'GW007'),
            ('Do you build online stores?', 'FAQ07'),
        ]:
            with self.subTest(question=question):
                self.assertIn(expected, {d.metadata['record_id'] for d in self.kb.retrieve(question)})

    def test_refresh_and_invalid_replacement(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'knowledge.docx'
            doc = Document(); doc.add_paragraph('GW001 Location'); doc.add_paragraph('Lucknow office'); doc.save(path)
            kb = KnowledgeBase(path)
            self.assertIn('Lucknow', kb.retrieve('office')[0].page_content)
            doc = Document(); doc.add_paragraph('GW001 Location'); doc.add_paragraph('New approved location'); doc.save(path)
            self.assertIn('New approved', kb.retrieve('location')[0].page_content)
            path.write_bytes(b'invalid document')
            with self.assertRaises(Exception): kb.retrieve('location')

    def test_duplicate_ids_are_rejected(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'knowledge.docx'
            doc = Document()
            for city in ['Lucknow', 'Other']:
                doc.add_paragraph('GW001 Location'); doc.add_paragraph(city)
            doc.save(path)
            with self.assertRaises(ValueError): read_records(path)

if __name__ == '__main__': unittest.main()
